import os
import shutil
import subprocess
import tempfile

import streamlit as st
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


# ---------- Replacement logic ----------

def replace_in_paragraph(paragraph, find, repl):
    """Replace `find` with `repl` in one paragraph, even when the text is split
    across several runs. Keeps the formatting of the run where the match starts."""
    count = 0
    start_at = 0
    while True:
        runs = paragraph.runs
        full = "".join(r.text for r in runs)
        idx = full.find(find, start_at)
        if idx == -1:
            break
        end = idx + len(find)

        pos = 0
        first = True
        for r in runs:
            r_start, r_end = pos, pos + len(r.text)
            pos = r_end
            if r_end <= idx or r_start >= end:
                continue  # run not touched by the match
            s = max(idx, r_start) - r_start
            e = min(end, r_end) - r_start
            if first:
                r.text = r.text[:s] + repl + r.text[e:]
                first = False
            else:
                r.text = r.text[:s] + r.text[e:]

        count += 1
        start_at = idx + len(repl)  # avoids endless loops if repl contains find
    return count


def replace_in_text_frame(text_frame, pairs):
    counts = [0] * len(pairs)
    for paragraph in text_frame.paragraphs:
        for i, (find, repl) in enumerate(pairs):
            counts[i] += replace_in_paragraph(paragraph, find, repl)
    return counts


def replace_in_shapes(shapes, pairs):
    counts = [0] * len(pairs)

    def add(new):
        for i, n in enumerate(new):
            counts[i] += n

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            add(replace_in_shapes(shape.shapes, pairs))
            continue
        if shape.has_text_frame:
            add(replace_in_text_frame(shape.text_frame, pairs))
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    add(replace_in_text_frame(cell.text_frame, pairs))
    return counts


def process_presentation(prs, pairs):
    totals = [0] * len(pairs)

    def add(new):
        for i, n in enumerate(new):
            totals[i] += n

    for slide in prs.slides:
        add(replace_in_shapes(slide.shapes, pairs))
        if slide.has_notes_slide:
            add(replace_in_text_frame(slide.notes_slide.notes_text_frame, pairs))

    # Master and layouts (footers, repeated headers, etc.)
    for master in prs.slide_masters:
        add(replace_in_shapes(master.shapes, pairs))
        for layout in master.slide_layouts:
            add(replace_in_shapes(layout.shapes, pairs))
    return totals


# ---------- PDF export ----------

def pptx_to_pdf(pptx_path, out_dir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice is not installed on this server (see packages.txt).")
    profile_dir = tempfile.mkdtemp(prefix="lo_profile_")
    cmd = [
        soffice,
        f"-env:UserInstallation=file://{profile_dir}",
        "--headless",
        "--convert-to", "pdf",
        "--outdir", out_dir,
        pptx_path,
    ]
    subprocess.run(cmd, check=True, timeout=240, capture_output=True)
    pdf_path = os.path.join(out_dir, os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf")
    if not os.path.exists(pdf_path):
        raise RuntimeError("PDF conversion failed.")
    return pdf_path


# ---------- App ----------

st.set_page_config(page_title="PPT Replace & PDF", page_icon="📄")
st.title("📄 PPT Find & Replace → PDF")

uploaded = st.file_uploader("Upload your PowerPoint (.pptx)", type=["pptx"])

st.subheader("Date")
c1, c2 = st.columns(2)
old_date = c1.text_input("Current date in the file", placeholder="e.g. 12 October 2026")
new_date = c2.text_input("New date", placeholder="e.g. 19 October 2026")

st.subheader("Opponent")
c3, c4 = st.columns(2)
old_opp = c3.text_input("Current opponent in the file", placeholder="e.g. Team A")
new_opp = c4.text_input("New opponent", placeholder="e.g. Team B")

if st.button("Replace and create PDF", type="primary", disabled=uploaded is None):
    pairs = []
    labels = []
    if old_date:
        pairs.append((old_date, new_date))
        labels.append("Date")
    if old_opp:
        pairs.append((old_opp, new_opp))
        labels.append("Opponent")

    if not pairs:
        st.error("Enter at least one 'current' value to search for.")
    else:
        with st.spinner("Working..."):
            try:
                with tempfile.TemporaryDirectory() as tmp:
                    base = os.path.splitext(uploaded.name)[0]
                    pptx_path = os.path.join(tmp, f"{base}.pptx")
                    with open(pptx_path, "wb") as f:
                        f.write(uploaded.getbuffer())

                    prs = Presentation(pptx_path)
                    counts = process_presentation(prs, pairs)
                    prs.save(pptx_path)

                    pdf_path = pptx_to_pdf(pptx_path, tmp)
                    with open(pdf_path, "rb") as f:
                        pdf_bytes = f.read()

                for label, n in zip(labels, counts):
                    if n:
                        st.success(f"{label}: replaced {n} occurrence(s)")
                    else:
                        st.warning(f"{label}: not found. Check the spelling and spacing.")

                st.download_button(
                    "⬇️ Download PDF",
                    data=pdf_bytes,
                    file_name=f"{base}.pdf",
                    mime="application/pdf",
                )
            except Exception as e:
                st.error(f"Something went wrong: {e}")
